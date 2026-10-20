import builtins
import importlib.util
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
APP_PATH = PROJECT_ROOT / "attendance_system_with_blink.py"
SPEC = importlib.util.spec_from_file_location("desktop_attendance_app", APP_PATH)
APP = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = APP
original_import = builtins.__import__


def import_without_optional_face_backends(name, *args, **kwargs):
    if name in {"face_recognition", "mediapipe"}:
        raise ImportError(f"{name} disabled for isolated desktop tests")
    return original_import(name, *args, **kwargs)


with patch("builtins.__import__", side_effect=import_without_optional_face_backends):
    SPEC.loader.exec_module(APP)


class DesktopAttendanceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)
        self.application = APP.QApplication.instance() or APP.QApplication([])

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_existing_database_users_are_preserved_and_get_unassigned_class(self):
        database_path = self.temp_path / "legacy.sqlite3"
        connection = sqlite3.connect(database_path)
        connection.execute(
            "CREATE TABLE users (student_id TEXT PRIMARY KEY, name TEXT NOT NULL, embedding BLOB NOT NULL)"
        )
        connection.execute(
            "CREATE TABLE attendance (id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "student_id TEXT NOT NULL, date DATE NOT NULL, first_seen TIME NOT NULL, "
            "last_seen TIME NOT NULL, FOREIGN KEY(student_id) REFERENCES users(student_id))"
        )
        connection.execute(
            "INSERT INTO users VALUES (?, ?, ?)",
            ("R-101", "Avery Student", b"existing-face-data"),
        )
        connection.execute(
            "INSERT INTO attendance (student_id, date, first_seen, last_seen) VALUES (?, ?, ?, ?)",
            ("R-101", "2026-10-19", "08:30:00", "08:30:00"),
        )
        connection.commit()
        connection.close()

        database = APP.Database(str(database_path))
        try:
            self.assertEqual(
                database.get_all_users(),
                [("R-101", "Avery Student", "Unassigned", "legacy")],
            )
            old_attendance = database.get_attendance("2026-10-19", "2026-10-19")
            self.assertEqual(len(old_attendance), 1)
            self.assertEqual(old_attendance[0][2], "Unassigned")
            database.update_user("R-101", class_name="MCA")
            self.assertEqual(
                database.get_all_users(),
                [("R-101", "Avery Student", "MCA", "legacy")],
            )
            class_name, attendance_date, attendance_time = database.mark_attendance("R-101")
            self.assertEqual(class_name, "MCA")
            self.assertEqual(attendance_date, APP.date.today().isoformat())
            self.assertRegex(attendance_time, r"^\d{2}:\d{2}:\d{2}$")
            attendance_rows = database.get_attendance(attendance_date, attendance_date)
            self.assertEqual(attendance_rows[0][2], "MCA")
        finally:
            database.conn.close()

    def test_four_uploaded_views_are_saved_as_sface_templates(self):
        database = APP.Database(str(self.temp_path / "photos.sqlite3"))
        widget = APP.RegistrationWidget(database)
        template = np.zeros(APP.SFACE_EMBEDDING_SIZE, dtype=np.float32)
        template[0] = 1.0

        class FakeRecognizer:
            def encode_enrollment_photo(self, image_path):
                self.last_path = image_path
                return template

        fake_recognizer = FakeRecognizer()
        widget.name_input.setText("Avery Student")
        widget.id_input.setText("R-101")
        widget.class_input.setText("MCA")
        widget.uploaded_photo_paths = [f"view-{number}.jpg" for number in range(4)]
        widget.face_recognizer = fake_recognizer
        with patch.object(APP.QMessageBox, "information"):
            widget.save_person()

        try:
            self.assertEqual(
                database.get_all_users(),
                [("R-101", "Avery Student", "MCA", "opencv_sface")],
            )
            _, templates = database.get_all_embeddings()["R-101"]
            self.assertEqual(templates.shape, (4, APP.SFACE_EMBEDDING_SIZE))
        finally:
            widget.stop_capture()
            database.conn.close()

    def test_manual_attendance_requires_three_unmatched_checks_and_known_roll_number(self):
        database = APP.Database(str(self.temp_path / "manual.sqlite3"))
        template = np.zeros((1, APP.SFACE_EMBEDDING_SIZE), dtype=np.float32)
        template[0, 0] = 1.0
        database.add_user("R-101", "Avery Student", "MCA", template)
        worker = APP.CameraThread(database)
        try:
            worker._update_unmatched_attempts(False, False, 1.0)
            self.assertEqual(worker._unmatched_attempts, 0)
            worker._update_unmatched_attempts(True, False, 1.0)
            worker._update_unmatched_attempts(True, False, 1.5)
            worker._update_unmatched_attempts(True, False, 2.0)
            worker._update_unmatched_attempts(True, False, 3.0)
            self.assertEqual(worker._unmatched_attempts, 3)
            worker._update_unmatched_attempts(True, True, 4.0)
            self.assertEqual(worker._unmatched_attempts, 0)
        finally:
            database.conn.close()

    def test_desktop_ui_exposes_student_and_attendance_actions(self):
        database = APP.Database(str(self.temp_path / "ui.sqlite3"))
        template = np.zeros((1, APP.SFACE_EMBEDDING_SIZE), dtype=np.float32)
        template[0, 0] = 1.0
        database.add_user("R-101", "Avery Student", "MCA", template)
        with patch.object(APP, "Database", return_value=database):
            window = APP.MainWindow()
        try:
            button_labels = {button.text() for button in window.findChildren(APP.QPushButton)}
            self.assertIn("Add new student", button_labels)
            self.assertIn("Take attendance", button_labels)
            self.assertEqual(window.registration_widget.class_input.placeholderText(), "e.g. MCA or MBA")
            live = window.live_widget
            live.on_unmatched_attempts_changed(3)
            self.assertTrue(live.manual_attendance_btn.isEnabled())
            with (
                patch.object(APP, "ATTENDANCE_RECORDS_DIR", self.temp_path / "records"),
                patch.object(APP.QInputDialog, "getText", return_value=("R-101", True)),
                patch.object(APP.QMessageBox, "information"),
            ):
                live.take_manual_attendance()
            self.assertEqual(
                database.get_attendance(APP.date.today().isoformat(), APP.date.today().isoformat())[0][:3],
                ("R-101", "Avery Student", "MCA"),
            )
        finally:
            window.close()
            self.application.processEvents()

    def test_attendance_workbook_is_date_based_class_wise_and_deduplicated(self):
        with patch.object(APP, "ATTENDANCE_RECORDS_DIR", self.temp_path):
            workbook_path = APP.write_daily_attendance_record(
                "R-101", "Avery Student", "MCA", "2026-10-20", "09:15:00"
            )
            APP.write_daily_attendance_record(
                "R-101", "Avery Student", "MCA", "2026-10-20", "09:15:00"
            )
            APP.write_daily_attendance_record(
                "R-102", "Riley Student", "MCA", "2026-10-20", "09:17:00"
            )
            APP.write_daily_attendance_record(
                "R-201", "Jordan Student", "MBA", "2026-10-20", "09:20:00"
            )

        self.assertEqual(workbook_path.name, "20-10-2026.xlsx")
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            self.assertEqual(set(workbook.sheetnames), {"MCA", "MBA"})
            mca_rows = list(workbook["MCA"].iter_rows(values_only=True))
            self.assertEqual(
                mca_rows[0],
                ("Name", "Roll Number", "Class", "Attendance Time"),
            )
            self.assertEqual(len(mca_rows), 3)
            self.assertEqual(
                mca_rows[1],
                ("Avery Student", "R-101", "MCA", "09:15:00"),
            )
            self.assertEqual(
                list(workbook["MBA"].iter_rows(values_only=True))[1],
                ("Jordan Student", "R-201", "MBA", "09:20:00"),
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
