import base64
from io import BytesIO
import re
import tempfile
import unittest
from datetime import date, time
import os
from pathlib import Path
from unittest.mock import patch

import numpy as np
import cv2
from openpyxl import load_workbook

from webapp import create_app, db
from webapp.models import Attendance, User
from webapp.vision import descriptor_backend


class WebApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        database_path = Path(self.temp_dir.name) / "test.sqlite3"
        self.app = create_app({
            "PROPAGATE_EXCEPTIONS": True,
            "SECRET_KEY": "unit-test-secret",
            "WEB_ADMIN_USERNAME": "admin",
            "WEB_ADMIN_PASSWORD": "test-password",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{database_path}",
        })
        self.client = self.app.test_client()
        with self.app.app_context():
            db.create_all()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()
            db.engine.dispose()
        self.temp_dir.cleanup()

    def authenticate(self):
        with self.client.session_transaction() as session:
            session["authenticated"] = True
            session["admin"] = "admin"
        response = self.client.get("/people")
        self.assertEqual(response.status_code, 200)

    def csrf_token(self, response):
        match = re.search(rb'<meta name="csrf-token" content="([^"]+)"', response.data)
        self.assertIsNotNone(match)
        return match.group(1).decode("ascii")

    def post_api(self, path, payload):
        page = self.client.get("/people")
        return self.client.post(
            path,
            json=payload,
            headers={"X-CSRFToken": self.csrf_token(page)},
        )

    def jpeg_data_url(self):
        image = np.zeros((160, 160, 3), dtype=np.uint8)
        image[40:120, 40:120] = 255
        success, encoded = cv2.imencode(".jpg", image)
        self.assertTrue(success)
        return "data:image/jpeg;base64," + base64.b64encode(encoded.tobytes()).decode("ascii")

    def test_unauthenticated_routes_redirect_and_health_is_public(self):
        self.assertEqual(self.client.get("/").status_code, 302)
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"status": "ok"})
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")
        self.assertIn("camera=(self)", response.headers["Permissions-Policy"])

    def test_login_uses_configured_credentials(self):
        csrf_token = self.csrf_token(self.client.get("/login"))
        response = self.client.post("/login", data={
            "username": "admin",
            "password": "wrong-password",
            "csrf_token": csrf_token,
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"incorrect", response.data)

        response = self.client.post("/login", data={
            "username": "admin",
            "password": "test-password",
            "csrf_token": csrf_token,
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.headers["Location"], "/")
        self.assertEqual(self.client.get("/").status_code, 200)

    def test_all_application_pages_render_for_authenticated_admin(self):
        self.authenticate()
        for path in ("/", "/live", "/people", "/analytics", "/reports"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_enrollment_and_recognition_store_descriptor_not_frame(self):
        self.authenticate()
        descriptor = np.zeros(1024, dtype=np.float64)
        frame_url = self.jpeg_data_url()
        with patch("webapp.routes.get_single_face_embedding", return_value=descriptor):
            response = self.post_api("/api/people", {
                "student_id": "S-001",
                "name": "Jordan Lee",
                "frames": [frame_url] * 6,
                "consent": True,
            })
        self.assertEqual(response.status_code, 201, response.json)

        with patch("webapp.routes.get_single_face_embedding", return_value=descriptor):
            first = self.post_api("/api/recognize", {"frame": frame_url})
            second = self.post_api("/api/recognize", {"frame": frame_url})
        self.assertEqual(first.status_code, 200, first.json)
        self.assertEqual(first.json["state"], "recognized")
        self.assertEqual(second.json["state"], "recognized")

        with self.app.app_context():
            person = db.session.get(User, "S-001")
            self.assertEqual(np.frombuffer(person.embedding, dtype=np.float64).shape, (1024,))
            records = db.session.query(Attendance).filter_by(student_id="S-001").all()
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].attendance_date, date.today())
            self.assertIsInstance(records[0].first_seen, time)

    def test_enrollment_requires_permission_and_unique_student_id(self):
        self.authenticate()
        payload = {
            "student_id": "S-002",
            "name": "Taylor Example",
            "frames": [self.jpeg_data_url()] * 3,
            "consent": False,
        }
        response = self.post_api("/api/people", payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("permission", response.json["error"])

        payload["consent"] = True
        descriptor = np.zeros(1024, dtype=np.float64)
        with patch("webapp.routes.get_single_face_embedding", return_value=descriptor):
            self.assertEqual(self.post_api("/api/people", payload).status_code, 201)
            duplicate = self.post_api("/api/people", payload)
        self.assertEqual(duplicate.status_code, 409)

    def test_report_download_contains_attendance_rows(self):
        self.authenticate()
        with self.app.app_context():
            person = User(student_id="S-003", name="Morgan Example", embedding=b"\0" * 8192)
            db.session.add(person)
            db.session.add(Attendance(
                student_id="S-003",
                attendance_date=date.today(),
                first_seen=time(9, 0),
                last_seen=time(9, 15),
            ))
            db.session.commit()

        response = self.client.get(f"/reports.xlsx?start={date.today()}&end={date.today()}")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Student ID", "Name", "Date", "First seen", "Last seen"))
        self.assertEqual(rows[1], ("S-003", "Morgan Example", date.today().isoformat(), "09:00:00", "09:15:00"))

    def test_csrf_is_required_for_authenticated_mutations(self):
        self.authenticate()
        with self.app.app_context():
            db.session.add(User(student_id="S-004", name="Delete Me", embedding=b"\0" * 8192))
            db.session.commit()
        response = self.client.post("/people/S-004/delete")
        self.assertEqual(response.status_code, 400)

    def test_create_app_requires_secrets_and_admin_configuration(self):
        with patch("webapp.load_dotenv"), patch.dict(
            os.environ,
            {"SECRET_KEY": "", "WEB_ADMIN_USERNAME": "", "WEB_ADMIN_PASSWORD": ""},
        ):
            with self.assertRaisesRegex(RuntimeError, "SECRET_KEY"):
                create_app({
                    "WEB_ADMIN_USERNAME": "admin",
                    "WEB_ADMIN_PASSWORD": "password",
                    "SQLALCHEMY_DATABASE_URI": "sqlite://",
                })
            with self.assertRaisesRegex(RuntimeError, "WEB_ADMIN_USERNAME"):
                create_app({
                    "SECRET_KEY": "test",
                    "WEB_ADMIN_USERNAME": "",
                    "WEB_ADMIN_PASSWORD": "",
                    "SQLALCHEMY_DATABASE_URI": "sqlite://",
                })

    def test_supabase_postgres_url_uses_psycopg_driver(self):
        supabase_url = (
            "postgresql://postgres.project-ref:db-password@"
            "aws-0-us-east-1.pooler.supabase.com:5432/postgres?sslmode=require"
        )
        with patch("webapp.load_dotenv"), patch.dict(
            os.environ,
            {
                "DATABASE_URL": supabase_url,
                "SECRET_KEY": "test-secret",
                "WEB_ADMIN_USERNAME": "admin",
                "WEB_ADMIN_PASSWORD": "test-password",
            },
        ):
            app = create_app()

        self.assertEqual(
            app.config["SQLALCHEMY_DATABASE_URI"],
            supabase_url.replace("postgresql://", "postgresql+psycopg://", 1),
        )

    def test_example_supabase_url_fails_with_configuration_error(self):
        with patch("webapp.load_dotenv"), patch.dict(
            os.environ,
            {
                "DATABASE_URL": (
                    "postgresql://postgres.<project-ref>:<password>"
                    "@<pooler-host>:5432/postgres?sslmode=require"
                ),
                "SECRET_KEY": "test-secret",
                "WEB_ADMIN_USERNAME": "admin",
                "WEB_ADMIN_PASSWORD": "test-password",
            },
        ):
            with self.assertRaisesRegex(RuntimeError, "example placeholders"):
                create_app()

    def test_configured_descriptor_backend_is_pixel_fallback(self):
        self.assertEqual(descriptor_backend(), "pixel_fallback")


if __name__ == "__main__":
    unittest.main()
