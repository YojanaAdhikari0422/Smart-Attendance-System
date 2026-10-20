"""Authenticated application pages and JSON endpoints."""

import base64
import binascii
import hmac
from datetime import date, datetime, timedelta
from io import BytesIO
from urllib.parse import urlsplit

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from werkzeug.exceptions import BadRequest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import db
from .models import Attendance, User
from .vision import ImageInputError, best_match, decode_frame, descriptor_backend, get_single_face_embedding


auth = Blueprint("auth", __name__)
main = Blueprint("main", __name__)
api = Blueprint("api", __name__)


def _parse_day(value, field_name):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        raise BadRequest(f"Choose a valid {field_name} date.") from None


def _get_report_dates():
    today = date.today()
    start_text = request.args.get("start", today.replace(day=1).isoformat())
    end_text = request.args.get("end", today.isoformat())
    start = _parse_day(start_text, "start")
    end = _parse_day(end_text, "end")
    if start > end:
        raise BadRequest("The start date must be on or before the end date.")
    return start, end


def _attendance_rows(start=None, end=None):
    query = select(Attendance).join(Attendance.user)
    if start is not None and end is not None:
        query = query.where(Attendance.attendance_date.between(start, end))
    return db.session.scalars(
        query.order_by(Attendance.attendance_date.desc(), Attendance.first_seen.desc())
    ).all()


def _decode_data_url(value):
    if not isinstance(value, str) or not value.startswith("data:image/jpeg;base64,"):
        raise ImageInputError("Camera frames must be JPEG images.")
    try:
        image_bytes = base64.b64decode(value.split(",", 1)[1], validate=True)
    except (binascii.Error, ValueError):
        raise ImageInputError("Could not decode the camera frame. Please capture it again.") from None
    return decode_frame(image_bytes)


def _mark_attendance(person):
    now = datetime.now()
    today = now.date()
    record = db.session.scalar(
        select(Attendance).where(
            Attendance.student_id == person.student_id,
            Attendance.attendance_date == today,
        )
    )
    if record is not None:
        last_seen = datetime.combine(today, record.last_seen)
        if now - last_seen < timedelta(seconds=15):
            return record, False
        record.last_seen = now.time().replace(microsecond=0)
        db.session.commit()
        return record, False

    record = Attendance(
        student_id=person.student_id,
        attendance_date=today,
        first_seen=now.time().replace(microsecond=0),
        last_seen=now.time().replace(microsecond=0),
    )
    db.session.add(record)
    try:
        db.session.commit()
        return record, True
    except IntegrityError:
        db.session.rollback()
        record = db.session.scalar(
            select(Attendance).where(
                Attendance.student_id == person.student_id,
                Attendance.attendance_date == today,
            )
        )
        if record is None:
            raise
        record.last_seen = now.time().replace(microsecond=0)
        db.session.commit()
        return record, False


@auth.route("/login", methods=["GET", "POST"])
def login():
    if session.get("authenticated"):
        return redirect(url_for("main.dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        expected_username = current_app.config["WEB_ADMIN_USERNAME"]
        expected_password = current_app.config["WEB_ADMIN_PASSWORD"]
        if hmac.compare_digest(username, expected_username) and hmac.compare_digest(password, expected_password):
            session.clear()
            session["authenticated"] = True
            session["admin"] = username
            session.permanent = True
            target = request.args.get("next", "")
            parsed_target = urlsplit(target)
            if (
                target.startswith("/")
                and not target.startswith("//")
                and "\\" not in target
                and not parsed_target.scheme
                and not parsed_target.netloc
            ):
                return redirect(target)
            return redirect(url_for("main.dashboard"))
        flash("The username or password is incorrect.", "error")
    return render_template("login.html", page_title="Sign in")


@auth.post("/logout")
def logout():
    session.clear()
    flash("You have been signed out.", "success")
    return redirect(url_for("auth.login"))


@main.get("/health")
def health():
    from sqlalchemy import text

    db.session.execute(text("SELECT 1"))
    return jsonify(status="ok")


@main.get("/")
def dashboard():
    today = date.today()
    total_people = db.session.scalar(select(func.count()).select_from(User)) or 0
    present_today = db.session.scalar(
        select(func.count()).select_from(Attendance).where(Attendance.attendance_date == today)
    ) or 0
    recent = _attendance_rows()[:8]
    rate = round(present_today / total_people * 100) if total_people else 0
    return render_template(
        "dashboard.html",
        page_title="Overview",
        total_people=total_people,
        present_today=present_today,
        attendance_rate=rate,
        recent=recent,
        today=today,
        active_page="overview",
    )


@main.get("/live")
def live():
    return render_template(
        "live.html",
        page_title="Live attendance",
        active_page="live",
        liveness_enabled=False,
        descriptor_backend=descriptor_backend(),
    )


@main.get("/people")
def people():
    search = request.args.get("q", "").strip()
    statement = select(User)
    if search:
        statement = statement.where(
            or_(User.student_id.ilike(f"%{search}%"), User.name.ilike(f"%{search}%"))
        )
    users = db.session.scalars(statement.order_by(User.name, User.student_id)).all()
    return render_template(
        "people.html",
        page_title="People",
        active_page="people",
        users=users,
        search=search,
        descriptor_backend=descriptor_backend(),
    )


@main.post("/people/<student_id>/delete")
def delete_person(student_id):
    person = db.session.get(User, student_id)
    if person is None:
        abort(404)
    try:
        db.session.delete(person)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        current_app.logger.exception("Failed to delete person %s", student_id)
        flash("Could not delete this person. Please try again.", "error")
        return redirect(url_for("main.people"))
    flash(f"Deleted {person.name} and their attendance history.", "success")
    return redirect(url_for("main.people"))


@main.post("/people/<student_id>/edit")
def edit_person(student_id):
    person = db.session.get(User, student_id)
    if person is None:
        abort(404)
    name = request.form.get("name", "").strip()
    if not name or len(name) > 160:
        flash("Enter a name between 1 and 160 characters.", "error")
        return redirect(url_for("main.people"))
    person.name = name
    try:
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        current_app.logger.exception("Failed to update person %s", student_id)
        flash("Could not update this person. Please try again.", "error")
        return redirect(url_for("main.people"))
    flash("Person updated.", "success")
    return redirect(url_for("main.people"))


@main.get("/analytics")
def analytics():
    start, end = _get_report_dates()
    rows = _attendance_rows(start, end)
    total_people = db.session.scalar(select(func.count()).select_from(User)) or 0
    present_days = len({row.student_id for row in rows})
    rate = round(present_days / total_people * 100) if total_people else 0
    return render_template(
        "analytics.html",
        page_title="Analytics",
        active_page="analytics",
        rows=rows,
        start=start,
        end=end,
        total_people=total_people,
        present_days=present_days,
        attendance_rate=rate,
    )


@main.get("/reports")
def reports():
    start, end = _get_report_dates()
    rows = _attendance_rows(start, end)
    return render_template(
        "reports.html",
        page_title="Reports",
        active_page="reports",
        rows=rows,
        start=start,
        end=end,
    )


@main.get("/reports.xlsx")
def export_report():
    start, end = _get_report_dates()
    rows = _attendance_rows(start, end)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    headers = ["Student ID", "Name", "Date", "First seen", "Last seen"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315FCE")
    for record in rows:
        values = [
            record.student_id,
            record.user.name if record.user else "Unknown",
            record.attendance_date.isoformat(),
            record.first_seen_text,
            record.last_seen_text,
        ]
        sheet.append([
            f"'{value}" if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            for value in values
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDE", (20, 30, 16, 16, 16)):
        sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"attendance_{start.isoformat()}_to_{end.isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@api.post("/people")
def create_person():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify(error="Send a JSON request with a name, student ID, and camera frames."), 400

    student_id = str(payload.get("student_id", "")).strip()
    name = str(payload.get("name", "")).strip()
    frames = payload.get("frames")
    consent = payload.get("consent") is True
    if not student_id or len(student_id) > 80 or not name or len(name) > 160:
        return jsonify(error="Enter a student ID (up to 80 characters) and name (up to 160 characters)."), 400
    if not consent:
        return jsonify(error="Confirm that you have permission to enroll and store this person's face descriptor."), 400
    if not isinstance(frames, list) or not 3 <= len(frames) <= 6:
        return jsonify(error="Capture between 3 and 6 face frames to enroll this person."), 400
    if db.session.get(User, student_id) is not None:
        return jsonify(error="That student ID is already registered."), 409

    try:
        embeddings = [get_single_face_embedding(_decode_data_url(frame)) for frame in frames]
    except ImageInputError as error:
        return jsonify(error=str(error)), 400
    mean_embedding = sum(embeddings) / len(embeddings)
    person = User(student_id=student_id, name=name, embedding=mean_embedding.astype("float64").tobytes())
    db.session.add(person)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify(error="That student ID is already registered."), 409
    except SQLAlchemyError:
        db.session.rollback()
        current_app.logger.exception("Failed to save a newly enrolled person")
        return jsonify(error="Could not save the person. Please try again."), 500
    return jsonify(message=f"{name} was enrolled successfully.", student_id=student_id), 201


@api.post("/recognize")
def recognize():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify(error="Send a camera frame as JSON."), 400
    try:
        frame = _decode_data_url(payload.get("frame"))
        embedding = get_single_face_embedding(frame)
    except ImageInputError as error:
        return jsonify(state="not_found", message=str(error)), 200

    people = db.session.scalars(select(User)).all()
    person, distance = best_match(embedding, people)
    if person is None:
        return jsonify(
            state="unknown",
            message="Face not recognized. Register the person first.",
        ), 200
    try:
        record, created = _mark_attendance(person)
    except SQLAlchemyError:
        db.session.rollback()
        current_app.logger.exception("Could not mark attendance for %s", person.student_id)
        return jsonify(error="Attendance could not be saved. Please try again."), 500
    return jsonify(
        state="recognized",
        student_id=person.student_id,
        name=person.name,
        message="Attendance recorded." if created else "Attendance updated.",
        first_seen=record.first_seen_text,
        distance=round(distance, 3) if distance is not None else None,
    )
